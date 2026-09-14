# ============================================================================
# Build the presentation deck from the rendered diagrams.
# ============================================================================
# docs/diagrams/*.png  ->  docs/diagrams/Westwood-Portal-Presentation.pptx
#
# Each diagram is 3840x2160 (16:9), so it is placed full-bleed on a
# 13.333x7.5in slide — exactly the same aspect, no cropping, 288 DPI.
#
# Usage:  py scripts/build-ppt.py
# ============================================================================

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

REPO = Path(__file__).resolve().parent.parent
DIAGRAMS = REPO / "docs" / "diagrams"
OUT = DIAGRAMS / "Westwood-Portal-Presentation.pptx"

# ── Brand palette (matches the diagrams) ────────────────────────────────────
NAVY = RGBColor.from_string("0B1B35")
NAVY_DARK = RGBColor.from_string("061120")
NAVY_LIGHT = RGBColor.from_string("122848")
GOLD = RGBColor.from_string("C49D3F")
GOLD_LIGHT = RGBColor.from_string("D8B86A")
OFFWHITE = RGBColor.from_string("FAF9F6")
CREAM = RGBColor.from_string("F5F2EC")
SLATE = RGBColor.from_string("7A8AA3")
SLATE_LIGHT = RGBColor.from_string("A3AFC2")
CHARCOAL = RGBColor.from_string("242C3D")
BORDER = RGBColor.from_string("E3DDD2")

SERIF = "Georgia"
SANS = "Segoe UI"

SLIDES = [
    ("01-system-flow.png", "System Flow", "Lead-to-matter lifecycle in eight steps",
     "Walks the full lifecycle: a visitor browses, submits an inquiry, the server validates and files it, "
     "admin reviews the intake queue and converts a qualified inquiry into a numbered matter, the lawyer and "
     "team are assigned, and the client collaborates in the portal until the matter closes with a complete audit trail."),
    ("02-use-case.png", "Use Cases", "Visitor · Client · Lawyer · Admin — who can do what",
     "Four actor lanes map every use case to the real permission model. Break-glass document access is admin-only "
     "and always audited. Clients message the team and exchange documents on their own matters; lawyers manage "
     "assigned matters, internal notes and the matter team."),
    ("03-erd.png", "Entity Relationship", "PostgreSQL data model with RLS on every table",
     "The schema as deployed: profiles with composite names, matters as the hub, matter_members / matter_notes / "
     "matter_events, documents with access levels and 15-minute break-glass grants, notifications and the "
     "audit_logs trail. Enum values are listed at the right."),
    ("04-flowchart.png", "Process Flowchart", "Decision points, validation loop, break-glass path",
     "Two decisions gate the intake: qualification by admin review, and the requirement that a registered client "
     "account exists before conversion. Invalid submissions loop back with a clear 400. The side panel shows the "
     "admin-only break-glass path for confidential documents."),
    ("05-dfd.png", "Data Flow", "Entities → processes → data stores (level 1)",
     "Level-1 data flow: visitors, clients, lawyers and admins feed four processes, which read and write seven "
     "data stores. Gold dashed lines are read paths (public pages, identity). Every other flow passes through "
     "row-level security."),
    ("06-system-architecture.png", "System Architecture", "React SPA → Supabase → PostgreSQL, end to end",
     "Five bands: users, the React 19 SPA hosted on Vercel, the Supabase platform (auth, API, realtime, storage, "
     "edge functions), the PostgreSQL data tier, and the cross-cutting security spine. The browser never holds "
     "anything more than the anon key."),
]

SECURITY_POINTS = [
    ("Row-level security on every table",
     "Policies are the only gate — client · lawyer · admin"),
    ("Break-glass is admin-only and audited",
     "Reason ≥ 15 characters · 15-minute expiry · lead lawyer notified"),
    ("The audit trail cannot be forged or erased",
     "Trigger-written vs client-reported rows are stamped · cleanup is admin-only with a 30-day floor"),
    ("Deactivated accounts are denied by the database",
     "The role helpers return nothing for inactive users — not just a UI state"),
    ("Nothing sensitive in the browser",
     "No service-role key, no custom server — one source of truth: PostgreSQL"),
]


def set_spacing(run, hundredths: int) -> None:
    """Letter-spacing (PPT `spc` attribute, in 1/100 pt)."""
    rpr = run._r.get_or_add_rPr()
    rpr.set("spc", str(hundredths))


def add_text(slide, x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP):
    """paragraphs: list of (runs, space_after_pt, line_spacing) where each run is
    (text, size_pt, bold, color, font, spacing_hundredths)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (runs, space_after, line_spacing) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if space_after is not None:
            p.space_after = Pt(space_after)
        if line_spacing is not None:
            p.line_spacing = line_spacing
        for text, size, bold, color, font, spc in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = font
            if spc:
                set_spacing(r, spc)
    return box


def add_rect(slide, x, y, w, h, fill, rounded=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if rounded:
        shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ── Slide 1 · Title ─────────────────────────────────────────────────────
    s = blank(prs)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    add_rect(s, 0.9, 2.02, 1.1, 0.055, GOLD)
    add_text(s, 0.9, 2.3, 11.5, 0.4, [([("WESTWOOD LAW FIRM PORTAL", 13, True, GOLD, SANS, 300)], None, None)])
    add_text(s, 0.9, 2.72, 11.5, 1.2, [
        ([("System Architecture & Process Flows", 40, True, RGBColor.from_string("FFFFFF"), SERIF, 0)], None, 1.05),
    ])
    add_text(s, 0.9, 4.06, 10.5, 0.9, [
        ([("System flow · use cases · data model · process flowchart · data flow · architecture",
           15, False, SLATE_LIGHT, SANS, 0)], None, 1.4),
    ])
    add_text(s, 0.9, 6.62, 11.5, 0.4, [
        ([("Westwood Law Firm — Client Portal", 11.5, False, SLATE, SANS, 0),
          ("      Prepared September 2026", 11.5, False, SLATE, SANS, 0)], None, None),
    ])
    set_notes(s, "Deck built from the six landscape diagrams, each rendered at 3840x2160 for crisp projection.")

    # ── Slide 2 · Contents ──────────────────────────────────────────────────
    s = blank(prs)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = OFFWHITE
    add_rect(s, 0.7, 0.62, 0.75, 0.05, GOLD)
    add_text(s, 0.7, 0.82, 11.9, 0.7, [
        ([("Contents", 30, True, NAVY, SERIF, 0)], None, None),
    ])
    cols_x = (0.7, 6.97)
    rows_y = (1.85, 3.55, 5.25)
    for i, (_, title, desc, _) in enumerate(SLIDES):
        x = cols_x[i % 2]
        y = rows_y[i // 2]
        add_rect(s, x, y, 5.65, 1.42, RGBColor.from_string("FFFFFF"), rounded=True, line=BORDER)
        add_text(s, x + 0.32, y + 0.24, 5.0, 1.0, [
            ([("%02d  " % (i + 1), 14, True, GOLD, SANS, 0),
              (title, 16, True, NAVY, SERIF, 0)], None, None),
            ([(desc, 11.5, False, SLATE, SANS, 0)], None, None),
        ])
    set_notes(s, "Six diagrams, in the order of the system: flow, use cases, data model, flowchart, data flow, architecture.")

    # ── Slides 3–8 · One full-bleed diagram each ────────────────────────────
    for png, title, _, note in SLIDES:
        s = blank(prs)
        s.shapes.add_picture(str(DIAGRAMS / png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        set_notes(s, note)

    # ── Slide 9 · Security at a glance ──────────────────────────────────────
    s = blank(prs)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY_DARK
    add_rect(s, 0.9, 0.72, 0.75, 0.05, GOLD)
    add_text(s, 0.9, 0.92, 11.5, 0.8, [
        ([("Security at a Glance", 30, True, RGBColor.from_string("FFFFFF"), SERIF, 0)], None, None),
    ])
    y = 2.1
    for head, sub in SECURITY_POINTS:
        add_rect(s, 0.9, y + 0.13, 0.1, 0.1, GOLD)
        add_text(s, 1.28, y - 0.06, 11.2, 0.8, [
            ([(head, 15, True, GOLD_LIGHT, SANS, 0)], None, None),
            ([(sub, 11.5, False, SLATE_LIGHT, SANS, 0)], None, None),
        ])
        y += 0.95
    set_notes(s, "Five properties that hold across the whole system — all enforced in the database, not just the UI.")

    prs.save(str(OUT))
    print(f"saved {OUT}  ({OUT.stat().st_size / 1024:.0f} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
